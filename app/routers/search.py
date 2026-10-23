"""Advanced search, Excel export, saved favorite filters."""
from __future__ import annotations

import io
import json
from urllib.parse import urlencode

from fastapi import APIRouter, Form, Query, Request
from fastapi.responses import RedirectResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app import config
from app.database import (
    build_search_where,
    db_session,
    get_grouping_options,
    lookups_by_type,
    supplier_tags_display,
    utc_now,
)
from app.deps import (
    current_user,
    flash,
    load_visible_columns,
    login_redirect,
    render,
    require_csrf,
    require_login,
)

router = APIRouter(tags=["search"])

FILTER_LABELS = {
    "q": "Quick / wildcard",
    "factory_trading_name": "Factory name",
    "contact_name": "Contact",
    "email": "Email",
    "city": "City",
    "website": "Website",
    "tel_whatsapp_wechat": "Tel / WhatsApp / WeChat",
    "notes": "Notes",
    "status": "Status",
    "interest_level": "Interest level",
    "meet": "Meet",
    "country": "Country",
    "cn_province": "CN Province",
    "nature_of_product": "Nature of Product",
    "item_group_name": "Item Group",
    "production_method": "Production method",
    "grouping": "Grouping",
    "tag": "Tag",
}

FILTER_KEYS = [
    "q",
    "factory_trading_name",
    "contact_name",
    "email",
    "city",
    "website",
    "tel_whatsapp_wechat",
    "notes",
    "status",
    "interest_level",
    "meet",
    "country",
    "cn_province",
    "nature_of_product",
    "item_group_name",
    "production_method",
    "grouping",
    "tag",
]


def _filters_from_request(request: Request) -> dict:
    out = {}
    for k in FILTER_KEYS:
        v = request.query_params.get(k, "")
        if v:
            out[k] = v
    return out



FILTER_TO_COLUMN = {
    "factory_trading_name": "factory_trading_name",
    "contact_name": "contact_name",
    "email": "email",
    "city": "city",
    "website": "website",
    "tel_whatsapp_wechat": "tel_whatsapp_wechat",
    "notes": "notes",
    "status": "status",
    "interest_level": "interest_level",
    "meet": "meet",
    "country": "country",
    "cn_province": "cn_province",
    "nature_of_product": "nature_of_product",
    "item_group_name": "item_group_name",
    "production_method": "production_method",
    "grouping": "tag_1",
    "tag": "tags",
}


def _visible_columns_with_filters(user_columns: list[str], filters: dict) -> list[str]:
    """Keep user column prefs, then append any columns used as active filters."""
    allowed = {k for k, _ in config.ALL_COLUMNS}
    out: list[str] = [c for c in user_columns if c in allowed]
    seen = set(out)
    for key, value in filters.items():
        if not value:
            continue
        col = FILTER_TO_COLUMN.get(key)
        if col and col in allowed and col not in seen:
            out.append(col)
            seen.add(col)
    return out


def _run_search(filters: dict) -> list[dict]:
    where, params = build_search_where(filters)
    with db_session() as conn:
        rows = conn.execute(
            f"SELECT * FROM suppliers WHERE {where} ORDER BY factory_trading_name COLLATE NOCASE",
            params,
        ).fetchall()
    return [{**dict(r), "tags": supplier_tags_display(dict(r))} for r in rows]


@router.get("/search")
def search_page(request: Request):
    user = current_user(request)
    if not user:
        return login_redirect()
    filters = _filters_from_request(request)
    results = _run_search(filters) if filters else []
    with db_session() as conn:
        lookups = lookups_by_type(conn)
        saved = []
        for r in conn.execute(
            "SELECT id, name, filters_json, created_at FROM saved_searches WHERE user_id = ? ORDER BY created_at DESC",
            (user["id"],),
        ).fetchall():
            d = dict(r)
            try:
                d["filters"] = json.loads(d["filters_json"])
            except Exception:
                d["filters"] = {}
            saved.append(d)
    visible = _visible_columns_with_filters(load_visible_columns(user["id"]), filters)
    active_filters = [
        {"key": k, "label": FILTER_LABELS.get(k, k), "value": filters[k]}
        for k in FILTER_KEYS
        if filters.get(k)
    ]
    return render(
        request,
        "search.html",
        {
            "filters": {k: filters.get(k, "") for k in FILTER_KEYS},
            "active_filters": active_filters,
            "results": results,
            "lookups": lookups,
            "status_values": config.STATUS_VALUES,
            "interest_level_values": config.INTEREST_LEVEL_VALUES,
            "grouping_options": get_grouping_options(),
            "saved_searches": saved,
            "searched": bool(filters),
            "max_saved": config.MAX_SAVED_SEARCHES,
            "visible_columns": visible,
            "column_labels": dict(config.ALL_COLUMNS),
            "all_columns": config.ALL_COLUMNS,
            "nav_active": "search",
        },
    )


def _cell_value_for_export(supplier: dict, key: str):
    """Match on-screen search result cell content for Excel."""
    if key == "tags":
        return supplier.get("tags") or ""
    if key == "notes":
        return supplier.get("notes") or ""
    return supplier.get(key) or ""


@router.get("/search/export")
def export_search(request: Request):
    """Excel export: same columns as search results, AutoFilter on, widths fitted."""
    user = current_user(request)
    if not user:
        return login_redirect()
    filters = _filters_from_request(request)
    results = _run_search(filters)
    visible = _visible_columns_with_filters(load_visible_columns(user["id"]), filters)
    labels = dict(config.ALL_COLUMNS)
    headers = [labels.get(k, k) for k in visible]

    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for s in results:
        ws.append([_cell_value_for_export(s, key) for key in visible])

    # AutoFilter on header row (ready for Excel column filters)
    last_row = max(1, len(results) + 1)
    last_col = max(1, len(headers))
    ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"

    # Freeze header
    ws.freeze_panes = "A2"

    # Column widths auto-sized to content
    for col_idx in range(1, last_col + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, max_row=last_row):
            val = row[0].value
            if val is None:
                continue
            max_len = max(max_len, len(str(val)))
        # Cap very wide columns; leave a little padding
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 48)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=suppliers_export.xlsx"},
    )


@router.post("/search/favorites")
def save_favorite(
    request: Request,
    csrf_token: str = Form(""),
    name: str = Form(...),
    filters_json: str = Form(...),
):
    user = require_login(request)
    require_csrf(request, csrf_token)
    name = name.strip()[:80]
    if not name:
        flash(request, "Favorite name is required.", "error")
        return RedirectResponse(url="/search", status_code=303)
    try:
        filters = json.loads(filters_json)
        if not isinstance(filters, dict):
            raise ValueError("invalid")
    except Exception:
        flash(request, "Invalid filter payload.", "error")
        return RedirectResponse(url="/search", status_code=303)
    # keep only known keys
    clean = {k: str(v) for k, v in filters.items() if k in FILTER_KEYS and v}
    with db_session() as conn:
        count = conn.execute(
            "SELECT COUNT(*) FROM saved_searches WHERE user_id = ?", (user["id"],)
        ).fetchone()[0]
        if count >= config.MAX_SAVED_SEARCHES:
            flash(
                request,
                f"You already have {config.MAX_SAVED_SEARCHES} saved searches. Delete one first.",
                "error",
            )
            return RedirectResponse(url="/search?" + urlencode(clean), status_code=303)
        conn.execute(
            "INSERT INTO saved_searches (user_id, name, filters_json, created_at) VALUES (?, ?, ?, ?)",
            (user["id"], name, json.dumps(clean), utc_now()),
        )
    flash(request, f'Saved search "{name}".', "success")
    return RedirectResponse(url="/search?" + urlencode(clean), status_code=303)


@router.post("/search/favorites/{fav_id}/delete")
def delete_favorite(request: Request, fav_id: int, csrf_token: str = Form("")):
    user = require_login(request)
    require_csrf(request, csrf_token)
    with db_session() as conn:
        conn.execute(
            "DELETE FROM saved_searches WHERE id = ? AND user_id = ?",
            (fav_id, user["id"]),
        )
    flash(request, "Saved search deleted.", "success")
    return RedirectResponse(url="/search", status_code=303)
